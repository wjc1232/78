# -*- coding: utf-8 -*-
"""
凉贸通 · 选品分析报告导出模块
支持导出：Excel / Word / PDF / PNG 图片
依赖：pip install openpyxl python-docx reportlab matplotlib
"""

import io
import re
import textwrap
from datetime import datetime

import pandas as pd

EUR_TO_CNY = 7.5


# ==================== 公共工具 ====================

def safe_filename(name: str) -> str:
    return re.sub(r'[\\/:*?"<>|\s]+', '_', str(name)).strip('_') or 'report'


def profit_rate(p: dict) -> float:
    revenue = float(p.get('price_eu', 0)) * EUR_TO_CNY
    cost = float(p.get('cost', 0))
    if revenue <= 0:
        return 0.0
    return (revenue - cost) / revenue * 100


def build_products_df(result: dict) -> pd.DataFrame:
    rows = []
    for i, p in enumerate(result.get('products', []) or [], 1):
        rows.append({
            '排名': f'TOP{i}',
            '产品名称': p.get('name', ''),
            '综合评分': p.get('score', ''),
            '采购成本(元)': p.get('cost', ''),
            '预计售价(€)': p.get('price_eu', ''),
            '利润率(%)': round(profit_rate(p), 1),
            '月销量预测(件)': p.get('sales_prediction', ''),
            '产品卖点': '；'.join(p.get('selling_points') or []),
            '风险提示': '；'.join(p.get('risks') or []),
        })
    return pd.DataFrame(rows)


def _display_width(s: str) -> int:
    """中文按 2 个字符宽度计算，用于 Excel 列宽"""
    return sum(2 if ord(ch) > 127 else 1 for ch in str(s))


# ==================== 1. Excel ====================

def _risk_level(text: str) -> str:
    """简易风险等级判断"""
    t = str(text)
    high_kw = ['认证', '合规', '罚款', '禁止', '侵权', '召回', '安全', '违规']
    mid_kw = ['竞争', '价格', '物流', '关税', '退货', '库存', '汇率', '季节性']
    for k in high_kw:
        if k in t:
            return '高'
    for k in mid_kw:
        if k in t:
            return '中'
    return '低'


def to_excel(result: dict, meta: dict) -> bytes:
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.chart import BarChart, LineChart, Reference

    products = result.get('products', []) or []
    trend = result.get('trend_data', []) or []

    # ---------- 1) 报告概览 ----------
    overview_rows = [{'项目': k, '内容': v} for k, v in meta.items()]
    overview_rows.append({'项目': '推荐产品数量', '内容': len(products)})
    if products:
        scores = [p.get('score', 0) for p in products]
        overview_rows.append({'项目': '平均综合评分', '内容': round(sum(scores) / len(scores), 1)})
        best = max(products, key=lambda p: p.get('score', 0))
        overview_rows.append({'项目': '最高评分产品', '内容': best.get('name', '')})
        avg_profit = sum(profit_rate(p) for p in products) / len(products)
        overview_rows.append({'项目': '平均利润率(%)', '内容': round(avg_profit, 1)})
        total_sales = sum(p.get('sales_prediction', 0) for p in products)
        overview_rows.append({'项目': '月销量预测合计(件)', '内容': total_sales})
        total_gross = sum(
            (p.get('price_eu', 0) * EUR_TO_CNY - p.get('cost', 0)) * p.get('sales_prediction', 0)
            for p in products
        )
        overview_rows.append({'项目': '月毛利润预测合计(元)', '内容': round(total_gross, 1)})
    overview_df = pd.DataFrame(overview_rows)

    # ---------- 2) 推荐产品 ----------
    prod_rows = []
    for i, p in enumerate(products, 1):
        pr = profit_rate(p)
        revenue_cny = p.get('price_eu', 0) * EUR_TO_CNY
        gross = revenue_cny - p.get('cost', 0)
        prod_rows.append({
            '排名': f'TOP{i}',
            '产品名称': p.get('name', ''),
            '综合评分': p.get('score', ''),
            '采购成本(元)': p.get('cost', ''),
            '预计售价(€)': p.get('price_eu', ''),
            '预计售价(元)': round(revenue_cny, 1),
            '单件毛利润(元)': round(gross, 1),
            '利润率(%)': round(pr, 1),
            '月销量预测(件)': p.get('sales_prediction', ''),
            '月毛利润(元)': round(gross * p.get('sales_prediction', 0), 1),
            '产品卖点': '；'.join(p.get('selling_points') or []),
            '风险提示': '；'.join(p.get('risks') or []),
        })
    prod_df = pd.DataFrame(prod_rows)

    # ---------- 3) 风险分析 ----------
    risk_rows = []
    for i, p in enumerate(products, 1):
        for r in (p.get('risks') or []):
            risk_rows.append({
                '产品': f"TOP{i} {p.get('name', '')}",
                '风险项': r,
                '风险等级': _risk_level(r),
            })
    risk_df = pd.DataFrame(risk_rows) if risk_rows else pd.DataFrame(
        {'提示': ['暂无风险项']})

    # ---------- 4) 市场趋势 ----------
    trend_df = pd.DataFrame(trend)
    if not trend_df.empty and '销量' in trend_df.columns:
        trend_df['环比增长率(%)'] = (
            trend_df['销量'].pct_change().mul(100).round(1).fillna(0))

    # ---------- 5) 选品建议 ----------
    sugg_df = pd.DataFrame({'选品建议': [result.get('suggestion', '')]})

    # ---------- 写入 Excel ----------
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine='openpyxl') as writer:
        overview_df.to_excel(writer, sheet_name='报告概览', index=False)
        prod_df.to_excel(writer, sheet_name='推荐产品', index=False)
        risk_df.to_excel(writer, sheet_name='风险分析', index=False)
        if not trend_df.empty:
            trend_df.to_excel(writer, sheet_name='市场趋势', index=False)
        sugg_df.to_excel(writer, sheet_name='选品建议', index=False)

        wb = writer.book

        # 通用样式
        header_fill = PatternFill('solid', fgColor='667EEA')
        header_font = Font(color='FFFFFF', bold=True, size=11)
        thin = Side(style='thin', color='D0D0D0')
        border = Border(left=thin, right=thin, top=thin, bottom=thin)

        for ws in wb.worksheets:
            for cell in ws[1]:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal='center', vertical='center')
            for row in ws.iter_rows():
                for cell in row:
                    cell.border = border
                    if cell.row > 1:
                        cell.alignment = Alignment(vertical='center', wrap_text=True)
            for col_idx in range(1, ws.max_column + 1):
                max_w = 10
                for row_idx in range(1, ws.max_row + 1):
                    v = ws.cell(row=row_idx, column=col_idx).value
                    if v is not None:
                        max_w = max(max_w, _display_width(v))
                ws.column_dimensions[get_column_letter(col_idx)].width = min(max_w + 4, 60)
            ws.row_dimensions[1].height = 24
            ws.freeze_panes = 'A2'

        # ---------- 「推荐产品」Sheet 内嵌图表 ----------
        ws = wb['推荐产品']
        n = len(prod_df)
        if n > 0:
            cats = Reference(ws, min_col=2, min_row=2, max_row=n + 1)

            chart1 = BarChart()
            chart1.type = 'bar'
            chart1.title = '产品综合评分'
            chart1.height = 8
            chart1.width = 16
            data1 = Reference(ws, min_col=3, min_row=1, max_row=n + 1)  # 综合评分列
            chart1.add_data(data1, titles_from_data=True)
            chart1.set_categories(cats)
            ws.add_chart(chart1, f"A{n + 4}")

            chart2 = BarChart()
            chart2.type = 'col'
            chart2.title = '各产品利润率(%)'
            chart2.height = 8
            chart2.width = 16
            data2 = Reference(ws, min_col=8, min_row=1, max_row=n + 1)  # 利润率列
            chart2.add_data(data2, titles_from_data=True)
            chart2.set_categories(cats)
            ws.add_chart(chart2, f"A{n + 22}")

        # ---------- 「市场趋势」Sheet 内嵌折线图 ----------
        if not trend_df.empty and '销量' in trend_df.columns:
            ws2 = wb['市场趋势']
            n2 = len(trend_df)
            line = LineChart()
            line.title = '销量趋势'
            line.y_axis.title = '销量'
            line.x_axis.title = '月份'
            line.height = 9
            line.width = 20
            data = Reference(ws2, min_col=2, min_row=1, max_row=n2 + 1)
            cats2 = Reference(ws2, min_col=1, min_row=2, max_row=n2 + 1)
            line.add_data(data, titles_from_data=True)
            line.set_categories(cats2)
            ws2.add_chart(line, f"A{n2 + 4}")

    buf.seek(0)
    return buf.getvalue()


# ==================== 2. Word ====================

def to_word(result: dict, meta: dict) -> bytes:
    from docx import Document
    from docx.shared import Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn

    CN = '微软雅黑'

    def style_run(run, size=10.5, bold=None, color=None):
        run.font.name = CN
        rpr = run._element.get_or_add_rPr()
        rfonts = rpr.get_or_add_rFonts()
        rfonts.set(qn('w:eastAsia'), CN)
        rfonts.set(qn('w:ascii'), CN)
        rfonts.set(qn('w:hAnsi'), CN)
        run.font.size = Pt(size)
        if bold is not None:
            run.font.bold = bold
        if color:
            run.font.color.rgb = RGBColor(*color)

    doc = Document()
    normal = doc.styles['Normal']
    normal.font.name = CN
    normal.font.size = Pt(10.5)
    normal.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), CN)

    # 大标题
    h = doc.add_heading(level=0)
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_run(h.add_run('凉贸通 · 选品分析报告'), size=22, bold=True, color=(0x4B, 0x4B, 0x8F))

    # 元信息
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_run(p.add_run('　|　'.join(f'{k}：{v}' for k, v in meta.items())),
              size=9, color=(0x88, 0x88, 0x88))

    # ---- 一、推荐产品表 ----
    h1 = doc.add_heading(level=1)
    style_run(h1.add_run('一、推荐产品 TOP5'), size=14, bold=True, color=(0x4B, 0x4B, 0x8F))

    df = build_products_df(result)
    cols = ['排名', '产品名称', '综合评分', '采购成本(元)', '预计售价(€)', '利润率(%)', '月销量预测(件)']
    table = doc.add_table(rows=1, cols=len(cols))
    table.style = 'Light Grid Accent 1'
    for i, c in enumerate(cols):
        cell = table.rows[0].cells[i]
        cell.text = ''
        style_run(cell.paragraphs[0].add_run(c), size=10, bold=True)
    for _, row in df.iterrows():
        cells = table.add_row().cells
        for i, c in enumerate(cols):
            cells[i].text = ''
            style_run(cells[i].paragraphs[0].add_run(str(row[c])), size=9.5)

    # ---- 二、产品详情 ----
    h1 = doc.add_heading(level=1)
    style_run(h1.add_run('二、产品详情与卖点'), size=14, bold=True, color=(0x4B, 0x4B, 0x8F))

    for i, prod in enumerate(result.get('products', []) or [], 1):
        h2 = doc.add_heading(level=2)
        style_run(h2.add_run(
            f"TOP{i}　{prod.get('name', '')}（综合评分 {prod.get('score', '-')} 分）"),
            size=12, bold=True, color=(0x33, 0x33, 0x33))

        if prod.get('selling_points'):
            pb = doc.add_paragraph()
            style_run(pb.add_run('核心卖点：'), size=10, bold=True)
            for sp in prod['selling_points']:
                bp = doc.add_paragraph(style='List Bullet')
                style_run(bp.add_run(str(sp)), size=10)

        if prod.get('risks'):
            pr = doc.add_paragraph()
            style_run(pr.add_run('风险提示：'), size=10, bold=True, color=(0xC0, 0x39, 0x2B))
            for rk in prod['risks']:
                bp = doc.add_paragraph(style='List Bullet')
                style_run(bp.add_run(str(rk)), size=10, color=(0xC0, 0x39, 0x2B))

    # ---- 三、市场趋势 ----
    trend = result.get('trend_data', []) or []
    if trend:
        h1 = doc.add_heading(level=1)
        style_run(h1.add_run('三、市场趋势数据'), size=14, bold=True, color=(0x4B, 0x4B, 0x8F))
        tdf = pd.DataFrame(trend)
        tt = doc.add_table(rows=1, cols=len(tdf.columns))
        tt.style = 'Light Grid Accent 1'
        for i, c in enumerate(tdf.columns):
            tt.rows[0].cells[i].text = ''
            style_run(tt.rows[0].cells[i].paragraphs[0].add_run(str(c)), size=10, bold=True)
        for _, row in tdf.iterrows():
            cells = tt.add_row().cells
            for i, c in enumerate(tdf.columns):
                cells[i].text = ''
                style_run(cells[i].paragraphs[0].add_run(str(row[c])), size=9.5)

    # ---- 四、选品建议 ----
    h1 = doc.add_heading(level=1)
    style_run(h1.add_run('四、选品建议'), size=14, bold=True, color=(0x4B, 0x4B, 0x8F))
    sp = doc.add_paragraph()
    style_run(sp.add_run(result.get('suggestion', '')), size=10.5)

    # 页脚
    fp = doc.add_paragraph()
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_run(fp.add_run(f"凉贸通 © 2026　|　生成时间：{datetime.now():%Y-%m-%d %H:%M}"),
              size=8, color=(0xAA, 0xAA, 0xAA))

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.getvalue()


# ==================== 3. PDF ====================

def to_pdf(result: dict, meta: dict) -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib import colors
    from reportlab.lib.units import mm
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle

    # 注册内置中文字体（无需额外字体文件）
    pdfmetrics.registerFont(UnicodeCIDFont('STSong-Light'))
    F = 'STSong-Light'

    base = getSampleStyleSheet()
    title_s = ParagraphStyle('t', parent=base['Title'], fontName=F, fontSize=20,
                             leading=26, textColor=colors.HexColor('#4B4B8F'))
    h1_s = ParagraphStyle('h1', parent=base['Heading1'], fontName=F, fontSize=13,
                          leading=18, spaceBefore=14, spaceAfter=6,
                          textColor=colors.HexColor('#4B4B8F'))
    h2_s = ParagraphStyle('h2', parent=base['Heading2'], fontName=F, fontSize=11,
                          leading=16, spaceBefore=10, spaceAfter=4,
                          textColor=colors.HexColor('#333333'))
    body_s = ParagraphStyle('b', parent=base['BodyText'], fontName=F, fontSize=9.5,
                            leading=15, textColor=colors.HexColor('#333333'))
    meta_s = ParagraphStyle('m', parent=body_s, fontSize=8.5,
                            textColor=colors.HexColor('#888888'))

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4,
                            leftMargin=18 * mm, rightMargin=18 * mm,
                            topMargin=16 * mm, bottomMargin=16 * mm,
                            title='凉贸通选品分析报告')
    story = []

    story.append(Paragraph('凉贸通 · 选品分析报告', title_s))
    story.append(Spacer(1, 4))
    story.append(Paragraph('　|　'.join(f'{k}：{v}' for k, v in meta.items()), meta_s))
    story.append(Spacer(1, 10))

    # ---- 产品表 ----
    story.append(Paragraph('一、推荐产品 TOP5', h1_s))
    head = ['排名', '产品名称', '评分', '成本(元)', '售价(€)', '利润率', '月销量']
    data = [[Paragraph(f'<font color="#FFFFFF"><b>{c}</b></font>', body_s) for c in head]]
    for i, p in enumerate(result.get('products', []) or [], 1):
        data.append([
            Paragraph(f'TOP{i}', body_s),
            Paragraph(str(p.get('name', '')), body_s),
            Paragraph(str(p.get('score', '')), body_s),
            Paragraph(str(p.get('cost', '')), body_s),
            Paragraph(str(p.get('price_eu', '')), body_s),
            Paragraph(f'{profit_rate(p):.1f}%', body_s),
            Paragraph(str(p.get('sales_prediction', '')), body_s),
        ])
    t = Table(data, colWidths=[15 * mm, 42 * mm, 15 * mm, 22 * mm, 20 * mm, 18 * mm, 24 * mm],
              repeatRows=1)
    t.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#667EEA')),
        ('GRID', (0, 0), (-1, -1), 0.4, colors.HexColor('#D0D0E0')),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F4F6FC')]),
        ('TOPPADDING', (0, 0), (-1, -1), 5),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
    ]))
    story.append(t)

    # ---- 产品详情 ----
    story.append(Paragraph('二、产品详情与卖点', h1_s))
    for i, p in enumerate(result.get('products', []) or [], 1):
        story.append(Paragraph(
            f"TOP{i}　{p.get('name', '')}（综合评分 {p.get('score', '-')} 分）", h2_s))
        for sp in (p.get('selling_points') or []):
            story.append(Paragraph(f'• {sp}', body_s))
        for rk in (p.get('risks') or []):
            story.append(Paragraph(
                f'<font color="#C0392B">⚠ {rk}</font>', body_s))
        story.append(Spacer(1, 4))

    # ---- 趋势数据 ----
    trend = result.get('trend_data', []) or []
    if trend:
        story.append(Paragraph('三、市场趋势数据', h1_s))
        tdf = pd.DataFrame(trend)
        tdata = [[Paragraph(f'<font color="#FFFFFF"><b>{c}</b></font>', body_s)
                  for c in tdf.columns]]
        for _, row in tdf.iterrows():
            tdata.append([Paragraph(str(row[c]), body_s) for c in tdf.columns])
        tt = Table(tdata, repeatRows=1)
        tt.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#667EEA')),
            ('GRID', (0, 0), (-1, -1), 0.4, colors.HexColor('#D0D0E0')),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F4F6FC')]),
        ]))
        story.append(tt)

    # ---- 建议 ----
    story.append(Paragraph('四、选品建议', h1_s))
    story.append(Paragraph(result.get('suggestion', ''), body_s))

    doc.build(story)
    buf.seek(0)
    return buf.getvalue()


# ==================== 4. PNG 图片 ====================

def _ensure_cjk_font():
    """
    返回一个可用的中文字体名；找不到返回 None。
    服务器已通过 packages.txt 安装 fonts-noto-cjk，
    因此这里主要就是把它找出来。
    """
    from matplotlib import font_manager

    # 优先：Noto CJK（云端已装）
    candidates = [
        'Noto Sans CJK SC', 'Noto Sans CJK JP',
        'Noto Sans CJK TC', 'Noto Serif CJK SC',
        # 本地常见
        'Microsoft YaHei', 'SimHei', 'SimSun',
        'PingFang SC', 'Hiragino Sans GB', 'Heiti SC',
        'Source Han Sans SC', 'WenQuanYi Micro Hei',
        'Arial Unicode MS',
    ]

    available = {f.name for f in font_manager.fontManager.ttflist}
    for name in candidates:
        if name in available:
            return name

    # 兜底：在系统字体目录里直接扫一遍 Noto CJK
    import glob
    for path in glob.glob('/usr/share/fonts/**/*CJK*.ttc', recursive=True) + \
                glob.glob('/usr/share/fonts/**/*CJK*.otf', recursive=True):
        try:
            font_manager.fontManager.addfont(path)
            return font_manager.FontProperties(fname=path).get_name()
        except Exception:
            continue

    return None


def to_png(result: dict, meta: dict) -> bytes:
    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt

    font = _find_cjk_font()
    if font:
        plt.rcParams['font.sans-serif'] = [font]
        plt.rcParams['font.family'] = 'sans-serif'
    plt.rcParams['axes.unicode_minus'] = False

    products = result.get('products', []) or []
    trend = result.get('trend_data', []) or []

    fig = plt.figure(figsize=(11, 14.5), dpi=140)
    fig.patch.set_facecolor('white')

    # ---- 标题 ----
    fig.text(0.05, 0.978, '凉贸通 · 选品分析报告', fontsize=21, fontweight='bold',
             color='#4B4B8F', ha='left', va='top')
    fig.text(0.05, 0.947, '　|　'.join(f'{k}：{v}' for k, v in meta.items()),
             fontsize=9.5, color='#888888', ha='left', va='top')
    fig.add_artist(plt.Line2D([0.05, 0.95], [0.933, 0.933], color='#667EEA', lw=2.5,
                              transform=fig.transFigure))

    # ---- 左：评分条形图 ----
    ax1 = fig.add_axes([0.06, 0.675, 0.40, 0.235])
    if products:
        names = [p.get('name', '') for p in products][::-1]
        scores = [p.get('score', 0) for p in products][::-1]
        bars = ax1.barh(names, scores, color='#667EEA', height=0.55)
        ax1.set_xlim(0, max(scores) * 1.25 if max(scores) else 100)
        for b, s in zip(bars, scores):
            ax1.text(b.get_width() * 1.02, b.get_y() + b.get_height() / 2,
                     f'{s}', va='center', fontsize=9, color='#444444')
        ax1.tick_params(labelsize=9)
        for sp in ['top', 'right']:
            ax1.spines[sp].set_visible(False)
        ax1.grid(axis='x', linestyle='--', alpha=0.3)
        ax1.set_axisbelow(True)
    ax1.set_title('产品综合评分', fontsize=12, fontweight='bold', color='#333333', pad=10)

    # ---- 右：趋势折线图 ----
    ax2 = fig.add_axes([0.56, 0.675, 0.39, 0.235])
    if trend:
        xs = [str(d.get('月份', '')) for d in trend]
        ys = [d.get('销量', 0) for d in trend]
        ax2.plot(range(len(xs)), ys, marker='o', color='#764BA2',
                 linewidth=2.2, markersize=5)
        ax2.set_xticks(range(len(xs)))
        ax2.set_xticklabels(xs, rotation=45, fontsize=8, ha='right')
        ax2.tick_params(axis='y', labelsize=9)
        ax2.grid(axis='y', linestyle='--', alpha=0.3)
        ax2.set_axisbelow(True)
        for sp in ['top', 'right']:
            ax2.spines[sp].set_visible(False)
    ax2.set_title('目标市场销量趋势', fontsize=12, fontweight='bold', color='#333333', pad=10)

    # ---- 产品表格 ----
    ax3 = fig.add_axes([0.03, 0.30, 0.94, 0.34])
    ax3.axis('off')
    if products:
        cols = ['排名', '产品名称', '评分', '成本(元)', '售价(€)', '利润率', '月销量']
        cell_text = []
        for i, p in enumerate(products, 1):
            cell_text.append([
                f'TOP{i}', str(p.get('name', '')), str(p.get('score', '')),
                str(p.get('cost', '')), str(p.get('price_eu', '')),
                f'{profit_rate(p):.1f}%', str(p.get('sales_prediction', '')),
            ])
        tbl = ax3.table(cellText=cell_text, colLabels=cols,
                        loc='center', cellLoc='center',
                        colWidths=[0.09, 0.30, 0.09, 0.13, 0.12, 0.11, 0.16])
        tbl.auto_set_font_size(False)
        tbl.set_fontsize(9.5)
        tbl.scale(1, 2.0)
        for j in range(len(cols)):
            tbl[0, j].set_facecolor('#667EEA')
            tbl[0, j].set_text_props(color='white', fontweight='bold')
        for i in range(1, len(cell_text) + 1):
            for j in range(len(cols)):
                if i % 2 == 0:
                    tbl[i, j].set_facecolor('#F2F4FB')
                tbl[i, j].set_edgecolor('#D8DCEA')
    ax3.set_title('推荐产品一览', fontsize=12, fontweight='bold',
                  color='#333333', pad=18)

    # ---- 选品建议 ----
    ax4 = fig.add_axes([0.05, 0.03, 0.90, 0.24])
    ax4.axis('off')
    ax4.text(0, 1.0, '选品建议', fontsize=12, fontweight='bold',
             color='#4B4B8F', va='top')
    sugg = result.get('suggestion', '') or ''
    wrapped = '\n'.join(textwrap.wrap(sugg, width=58)) or '（无）'
    ax4.text(0, 0.80, wrapped, fontsize=10.5, va='top',
             linespacing=1.9, color='#333333')

    buf = io.BytesIO()
    fig.savefig(buf, format='png', facecolor='white', bbox_inches='tight')
    plt.close(fig)
    buf.seek(0)
    return buf.getvalue()




# ==================== 统一入口 ====================

EXT = {'Excel': '.xlsx', 'Word': '.docx', 'PDF': '.pdf', '图片': '.png'}
MIME = {
    'Excel': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    'Word': 'application/vnd.openxmlformats-officedocument.wordprocessingml.document',
    'PDF': 'application/pdf',
    '图片': 'image/png',
}


def export_selection_report(result: dict, meta: dict, fmt: str = 'Excel',
                            name_hint: str = ''):
    """返回 (bytes, 文件名, MIME 类型)"""
    base = (f"选品分析报告_{safe_filename(name_hint)}_"
            f"{datetime.now():%Y%m%d}")

    if fmt == 'Excel':
        data = to_excel(result, meta)
    elif fmt == 'Word':
        data = to_word(result, meta)
    elif fmt == 'PDF':
        data = to_pdf(result, meta)
    elif fmt == '图片':
        data = to_png(result, meta)
    else:
        raise ValueError(f'不支持的导出格式：{fmt}')

    return data, base + EXT[fmt], MIME[fmt]
